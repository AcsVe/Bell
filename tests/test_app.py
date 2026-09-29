import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from app import create_app  # noqa: E402
from models import (db, Device, Grade, GradePeriodTime, Period, Section, Stage,  # noqa: E402
                    StageDayPeriodTime, Teacher, TeacherSection)


@pytest.fixture
def app(tmp_path):
    app = create_app({'TESTING': True,
                      'SQLALCHEMY_DATABASE_URI': f'sqlite:///{tmp_path}/t.db'})
    yield app


@pytest.fixture
def admin(app):
    c = app.test_client()
    r = c.post('/admin/login', data={'username': 'admin', 'password': 'admin'})
    assert r.status_code == 302
    return c


def _setup_school(admin, app):
    admin.post('/admin/stages/add', data={'name_ar': 'المرحلة الأساسية 1-3', 'name_en': 'Primary 1-3'})
    with app.app_context():
        st = Stage.query.one()
        sid = st.id
        assert sorted(st.weekday_list) == [0, 1, 2, 3, 6]
    for kind, dur in [('class', 45), ('class', 45), ('break', 20), ('class', 45)]:
        admin.post(f'/admin/stages/{sid}/periods/add', data={'kind': kind, 'duration': dur})
    admin.post(f'/admin/stages/{sid}/grades/add', data={'names': 'الصف الأول، الصف الثاني'})
    with app.app_context():
        g1 = Grade.query.filter_by(name_ar='الصف الأول').one()
        gid = g1.id
    admin.post(f'/admin/grades/{gid}/sections/add', data={'names': 'أ، ب'})
    admin.post('/admin/teachers/add', data={'names': 'أحمد\nسارة\n'})
    return sid, gid


def test_admin_crud_and_labels(admin, app):
    sid, gid = _setup_school(admin, app)
    with app.app_context():
        labels = [(p.kind, p.label_ar, p.label_en) for p in Period.query.order_by(Period.number)]
        assert labels == [('class', 'الحصة الأولى', 'Period 1'), ('class', 'الحصة الثانية', 'Period 2'),
                          ('break', 'الاستراحة', 'Break'), ('class', 'الحصة الثالثة', 'Period 3')]
        secs = Section.query.all()
        assert len(secs) == 2 and all(len(s.code) == 6 for s in secs)
        ts = Teacher.query.all()
        codes = {s.code for s in secs} | {t.code for t in ts}
        assert len(codes) == 4   # unique across both tables
    # every admin page renders in both languages
    for lang in ('ar', 'en'):
        admin.get(f'/lang/{lang}')
        for url in ['/admin/', f'/admin/stages/{sid}', f'/admin/grades/{gid}', '/admin/teachers',
                    '/admin/codes', '/admin/settings', f'/admin/stages/{sid}?day=3',
                    f'/admin/grades/{gid}?day=3']:
            r = admin.get(url)
            assert r.status_code == 200, url
        with app.app_context():
            s1 = Section.query.first().id
            t1 = Teacher.query.first().id
        assert admin.get(f'/admin/sections/{s1}').status_code == 200
        assert admin.get(f'/admin/teachers/{t1}').status_code == 200


def test_teacher_dropdown_inactive_after_link(admin, app):
    sid, gid = _setup_school(admin, app)
    with app.app_context():
        s1, s2 = [s.id for s in Section.query.order_by(Section.id)]
        t = Teacher.query.filter_by(name_ar='أحمد').one().id
    admin.post(f'/admin/sections/{s1}/teachers/add', data={'teacher_id': t})
    admin.post(f'/admin/sections/{s1}/teachers/add', data={'teacher_id': t})   # duplicate ignored
    html = admin.get(f'/admin/sections/{s1}').get_data(as_text=True)
    assert f'value="{t}" disabled' in html
    html2 = admin.get(f'/admin/sections/{s2}').get_data(as_text=True)
    assert f'value="{t}" disabled' not in html2 and f'value="{t}"' in html2
    with app.app_context():
        assert TeacherSection.query.count() == 1


def _times(rows):
    return [(r['start'], r['end']) for r in rows]


def test_chained_times_and_overrides(admin, app):
    sid, gid = _setup_school(admin, app)
    from services import resolve_grade_day
    with app.app_context():
        g = db.session.get(Grade, gid)
        assert _times(resolve_grade_day(g, 6)) == [('08:00', '08:45'), ('08:45', '09:30'),
                                                   ('09:30', '09:50'), ('09:50', '10:35')]
        pids = [p.id for p in Period.query.order_by(Period.number)]
    # Thursday: whole stage has 40-minute classes
    admin.post(f'/admin/stages/{sid}/day-times', data={'weekday': 3, 'day_start': '08:00',
               'n1_dur': 40, 'n2_dur': 40, 'n3_dur': 20, 'n4_dur': 40})
    # Grade 1 on Thursday starts later
    admin.post(f'/admin/grades/{gid}/times', data={'weekday': 3, 'day_start': '08:30',
               'n1_dur': 40, 'n2_dur': 40, 'n3_dur': 20, 'n4_dur': 40})
    with app.app_context():
        g = db.session.get(Grade, gid)
        g2 = Grade.query.filter_by(name_ar='الصف الثاني').one()
        assert _times(resolve_grade_day(g2, 3)) == [('08:00', '08:40'), ('08:40', '09:20'),
                                                    ('09:20', '09:40'), ('09:40', '10:20')]
        thu = resolve_grade_day(g, 3)
        assert thu[0]['source'] == 'grade' and thu[-1]['end'] == '10:50'
        assert resolve_grade_day(g, 6)[0]['source'] == 'default'
    # changing the default duration of period 1 shifts every later default period
    form = {'day_start': '07:45'}
    for pid, d in zip(pids, [50, 45, 20, 45]):
        form.update({f'p{pid}_duration': d, f'p{pid}_kind': 'break' if d == 20 else 'class'})
    admin.post(f'/admin/stages/{sid}/periods/save', data=form)
    with app.app_context():
        g = db.session.get(Grade, gid)
        assert _times(resolve_grade_day(g, 6))[-1] == ('09:40', '10:25')
        assert resolve_grade_day(g, 3)[-1]['end'] == '10:50'     # override untouched
    # adding a period extends every override set with the default duration, gap-free
    admin.post(f'/admin/stages/{sid}/periods/add', data={'kind': 'class', 'duration': 45})
    with app.app_context():
        g = db.session.get(Grade, gid)
        thu = resolve_grade_day(g, 3)
        assert thu[-1]['start'] == '10:50' and thu[-1]['end'] == '11:35' and thu[-1]['source'] == 'grade'
        for wd in (3, 6):
            rows = resolve_grade_day(g, wd)
            assert all(a['end'] == b['start'] for a, b in zip(rows, rows[1:]))
    # deleting the break moves later periods up, in defaults and overrides
    admin.post(f'/admin/periods/{pids[2]}/delete')
    with app.app_context():
        g = db.session.get(Grade, gid)
        for wd in (3, 6):
            rows = resolve_grade_day(g, wd)
            assert len(rows) == 4 and all(a['end'] == b['start'] for a, b in zip(rows, rows[1:]))
        assert resolve_grade_day(g, 3)[2]['start'] == '09:50'
    # reset returns grade 1 to the stage-day timing; all_days writes every school day
    admin.post(f'/admin/grades/{gid}/times', data={'weekday': 3, 'reset': '1'})
    with app.app_context():
        assert resolve_grade_day(db.session.get(Grade, gid), 3)[0]['source'] == 'stage'
    admin.post(f'/admin/grades/{gid}/times', data={'weekday': 6, 'all_days': '1', 'day_start': '08:00',
               'n1_dur': 30, 'n2_dur': 30, 'n4_dur': 30, 'n5_dur': 30})
    with app.app_context():
        assert GradePeriodTime.query.filter_by(grade_id=gid).count() == 5 * 4
    # invalid input is rejected without changes
    admin.post(f'/admin/stages/{sid}/day-times', data={'weekday': 0, 'day_start': '23:00',
               'n1_dur': 90, 'n2_dur': 90, 'n4_dur': 90, 'n5_dur': 90})
    with app.app_context():
        assert StageDayPeriodTime.query.filter_by(weekday=0).count() == 0


def test_code_login_schedule_and_versioning(admin, app):
    sid, gid = _setup_school(admin, app)
    with app.app_context():
        sec = Section.query.order_by(Section.id).first()
        s_code, s_id = sec.code, sec.id
        t = Teacher.query.filter_by(name_ar='سارة').one()
        t_code, t_id = t.code, t.id
        s_ids = [s.id for s in Section.query.all()]
    for s in s_ids:
        admin.post(f'/admin/sections/{s}/teachers/add', data={'teacher_id': t_id})

    dev = app.test_client()
    assert dev.get('/api/schedule').status_code == 401
    r = dev.post('/login', data={'code': 'zzzzzz'})
    assert r.status_code == 200 and 'data-error="invalid"' in r.get_data(as_text=True)
    r = dev.post('/login', data={'code': ' ' + s_code.lower()[:3] + '-' + s_code.lower()[3:]})
    assert r.status_code == 302 and r.headers['Location'].endswith('/app')
    r = dev.get('/api/schedule', headers={'X-Device-Id': 'dev-1'})
    data = r.get_json()
    assert data['who']['type'] == 'section'
    assert len(data['targets']) == 1 and data['targets'][0]['sectionId'] == s_id
    day = data['days'][str(gid)]['6']
    assert [p['ar'] for p in day] == ['الحصة الأولى', 'الحصة الثانية', 'الاستراحة', 'الحصة الثالثة']
    assert data['tone'].endswith('chime.wav')
    v, k = data['version'], data['key']
    assert k == f'section:{s_id}'
    assert dev.get(f'/api/schedule?v={v}&k={k}').get_json() == {'unchanged': True, 'version': v}
    # same version but a different identity (new code on the device) → full payload
    assert 'unchanged' not in dev.get(f'/api/schedule?v={v}&k=teacher:1').get_json()
    admin.post(f'/admin/stages/{sid}/periods/save', data={})   # any change bumps version
    assert 'unchanged' not in dev.get(f'/api/schedule?v={v}&k={k}').get_json()
    with app.app_context():
        assert Device.query.filter_by(device_uid='dev-1', code_type='section', owner_id=s_id).count() == 1

    # teacher code covers all linked sections
    tdev = app.test_client()
    tdev.post('/login', data={'code': t_code})
    td = tdev.get('/api/schedule').get_json()
    assert td['who']['type'] == 'teacher' and len(td['targets']) == 2

    # regenerating the code invalidates the old session
    admin.post(f'/admin/sections/{s_id}/regen')
    assert dev.get('/api/schedule').status_code == 401
    with app.app_context():
        assert Device.query.filter_by(owner_id=s_id, code_type='section').count() == 0

    # disabled teacher is rejected
    admin.post(f'/admin/teachers/{t_id}/update', data={'name_ar': 'سارة'})
    assert tdev.get('/api/schedule').status_code == 401


def test_media_and_tone_upload(admin, app):
    sid, _ = _setup_school(admin, app)
    png = b'\x89PNG\r\n\x1a\n' + b'0' * 20
    admin.post(f'/admin/stages/{sid}/media', data={'logo': (io.BytesIO(png), 'logo.png', 'image/png')},
               content_type='multipart/form-data')
    with app.app_context():
        url = db.session.get(Stage, sid).logo.url
    r = admin.get(url)
    assert r.status_code == 200 and r.data == png
    # wrong type rejected
    admin.post(f'/admin/stages/{sid}/media', data={'background': (io.BytesIO(b'x'), 'a.txt', 'text/plain')},
               content_type='multipart/form-data')
    with app.app_context():
        assert db.session.get(Stage, sid).background is None
    admin.post('/admin/settings', data={'tone': (io.BytesIO(b'RIFF....'), 'bell.wav', 'audio/wav')},
               content_type='multipart/form-data')
    from models import AppSetting
    with app.app_context():
        assert AppSetting.get('tone_id')
    admin.post(f'/admin/stages/{sid}/media', data={'remove_logo': '1'}, content_type='multipart/form-data')
    with app.app_context():
        assert db.session.get(Stage, sid).logo is None


def test_delete_cascades(admin, app):
    sid, gid = _setup_school(admin, app)
    with app.app_context():
        s1 = Section.query.first().id
        t = Teacher.query.first().id
    admin.post(f'/admin/sections/{s1}/teachers/add', data={'teacher_id': t})
    admin.post(f'/admin/stages/{sid}/delete')
    with app.app_context():
        assert Stage.query.count() == 0 and Section.query.count() == 0
        assert TeacherSection.query.count() == 0 and Teacher.query.count() == 2


def test_pwa_routes(app):
    c = app.test_client()
    r = c.get('/sw.js')
    assert r.status_code == 200 and 'no-cache' in r.headers['Cache-Control']
    assert c.get('/manifest.webmanifest').status_code == 200
    assert c.get('/app').status_code == 200
    assert c.get('/login').status_code == 200
    assert c.get('/static/sounds/chime.wav').status_code == 200


def test_native_header_auth_and_cors(admin, app):
    _setup_school(admin, app)
    with app.app_context():
        code = Section.query.first().code
    c = app.test_client()   # no cookies at all, like a native app
    r = c.options('/api/schedule', headers={'Origin': 'capacitor://localhost',
                  'Access-Control-Request-Headers': 'x-bell-code'})
    assert r.status_code == 204 and r.headers['Access-Control-Allow-Origin'] == '*'
    assert 'X-Bell-Code' in r.headers['Access-Control-Allow-Headers']
    assert c.post('/api/login', json={'code': 'nope'}).status_code == 401
    r = c.post('/api/login', json={'code': code.lower()})
    assert r.status_code == 200 and r.get_json()['code'] == code and r.get_json()['type'] == 'section'
    r = c.get('/api/schedule', headers={'X-Bell-Code': code, 'X-Device-Id': 'android-1'})
    assert r.status_code == 200 and r.get_json()['who']['type'] == 'section'
    assert r.headers['Access-Control-Allow-Origin'] == '*'
    assert c.get('/api/schedule', headers={'X-Bell-Code': 'WRONG1'}).status_code == 401
    # admin pages never get CORS headers
    assert 'Access-Control-Allow-Origin' not in admin.get('/admin/').headers


def test_normalize_db_url():
    from app import normalize_db_url
    for raw in ('postgres://u:p@h/db', 'postgresql://u:p@h/db', 'postgresql+psycopg2://u:p@h/db',
                'postgresql+psycopg://u:p@h/db', '  postgresql://u:p@h/db  '):
        assert normalize_db_url(raw) == 'postgresql+psycopg://u:p@h/db'
    assert normalize_db_url('sqlite:///x.db') == 'sqlite:///x.db'


def _fill(xlsx_bytes, rows):
    """rows: list of (name_ar, name_en, [section header names to mark])"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb['المعلمون']
    headers = [c.value for c in ws[2]]
    for r in range(3, ws.max_row + 1):          # clear pre-filled rows
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).value = None
    for i, (ar, en, marks) in enumerate(rows):
        ws.cell(row=3 + i, column=1).value = ar
        ws.cell(row=3 + i, column=2).value = en
        for m in marks:
            ws.cell(row=3 + i, column=headers.index(m) + 1).value = '✓'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _upload(admin, sid, data, sync=True):
    form = {'stage_id': sid, 'file': (io.BytesIO(data), 't.xlsx')}
    if sync:
        form['sync'] = '1'
    return admin.post('/admin/teachers/import', data=form, content_type='multipart/form-data',
                      follow_redirects=True).get_data(as_text=True)


def test_teacher_template_import_by_stage(admin, app):
    import openpyxl
    sid, gid = _setup_school(admin, app)
    # second stage for a shared teacher
    admin.post('/admin/stages/add', data={'name_ar': 'المرحلة الثانوية'})
    with app.app_context():
        sid2 = Stage.query.filter_by(name_ar='المرحلة الثانوية').one().id
    admin.post(f'/admin/stages/{sid2}/grades/add', data={'names': 'الصف العاشر'})
    with app.app_context():
        g10 = db.session.get(Stage, sid2).grades[0].id
    admin.post(f'/admin/grades/{g10}/sections/add', data={'names': 'أ'})

    r = admin.get(f'/admin/teachers/template/{sid}')
    assert r.status_code == 200 and r.mimetype.endswith('sheet')
    ws = openpyxl.load_workbook(io.BytesIO(r.data))['المعلمون']
    assert ws['A1'].value == f'stage:{sid}' and ws.row_dimensions[1].hidden
    assert [c.value for c in ws[2]][5:] == ['الصف الأول - أ', 'الصف الأول - ب']

    data = _fill(r.data, [('منى  الخطيب', 'Muna', ['الصف الأول - أ', 'الصف الأول - ب']),
                          ('رامي عودة', '', ['الصف الأول - ب']),
                          ('أحمد', '', [])])   # existing teacher from _setup_school, no links
    html = _upload(admin, sid, data)
    assert 'جديد 2' in html and 'ربط 3' in html
    with app.app_context():
        muna = Teacher.query.filter_by(name_ar='منى الخطيب').one()   # whitespace normalized
        assert muna.name_en == 'Muna' and len(muna.code) == 6
        assert sorted(s.name_ar for s in muna.sections) == ['أ', 'ب']
        assert Teacher.query.filter_by(name_ar='أحمد').count() == 1  # matched, not duplicated
        muna_code = muna.code

    # re-download is pre-filled with current links and codes
    ws = openpyxl.load_workbook(io.BytesIO(admin.get(f'/admin/teachers/template/{sid}').data))['المعلمون']
    rows = {ws.cell(row=r, column=1).value: [ws.cell(row=r, column=c).value for c in (4, 6, 7)]
            for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value}
    assert rows['منى الخطيب'] == [muna_code, '✓', '✓'] and rows['رامي عودة'][1:] == [None, '✓']

    # same teacher in the other stage's template keeps one code
    t2 = admin.get(f'/admin/teachers/template/{sid2}').data
    _upload(admin, sid2, _fill(t2, [('منى الخطيب', '', ['الصف العاشر - أ'])]))
    with app.app_context():
        muna = Teacher.query.filter_by(name_ar='منى الخطيب').one()
        assert muna.code == muna_code and len(muna.sections) == 3

    # sync: unmarking ب in stage 1 removes only that link (stage-2 link untouched)
    t1 = admin.get(f'/admin/teachers/template/{sid}').data
    _upload(admin, sid, _fill(t1, [('منى الخطيب', 'Muna', ['الصف الأول - أ'])]))
    with app.app_context():
        muna = Teacher.query.filter_by(name_ar='منى الخطيب').one()
        assert sorted(s.full_name('ar') for s in muna.sections) == ['الصف الأول - أ', 'الصف العاشر - أ']
        # رامي not in the file → untouched
        assert len(Teacher.query.filter_by(name_ar='رامي عودة').one().sections) == 1

    # without sync nothing is removed
    _upload(admin, sid, _fill(t1, [('منى الخطيب', '', [])]), sync=False)
    with app.app_context():
        assert len(Teacher.query.filter_by(name_ar='منى الخطيب').one().sections) == 2

    # wrong stage is refused, nothing changes
    html = _upload(admin, sid2, _fill(t1, [('جديد', '', ['الصف الأول - أ'])]))
    assert 'مرحلة أخرى' in html
    with app.app_context():
        assert Teacher.query.filter_by(name_ar='جديد').count() == 0

    # a hand-made sheet without the hidden key row still works via header names
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['الاسم', 'English', 'الكود', 'الصف الأول - ب', 'عمود غريب'])
    ws.append(['ليلى', '', '', 'x', 'x'])
    buf = io.BytesIO()
    wb.save(buf)
    html = _upload(admin, sid, buf.getvalue())
    assert 'عمود غريب' in html
    with app.app_context():
        assert [s.name_ar for s in Teacher.query.filter_by(name_ar='ليلى').one().sections] == ['ب']

    # not an Excel file
    html = _upload(admin, sid, b'not excel')
    assert 'ليس ملف Excel' in html

    # codes export and page render
    r = admin.get(f'/admin/teachers/codes/{sid}')
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.data))['أكواد المعلمين']
    assert any(row[1] == muna_code for row in ws.iter_rows(values_only=True))
    assert 'تحميل القالب' in admin.get('/admin/teachers').get_data(as_text=True)


CSV_SAMPLE = ('\ufeff"Name","PrimarySmtpAddress","RecipientType"\n'
              '"amira.test","Amira.Test@school.example","UserMailbox"\n'
              '"esra\'a.thyab","esraa.t@school.example","UserMailbox"\n'
              '"22eacba1-2740-4d98-a273-3332c7d9fb8b","Rasha.AlAyyad@school.example","UserMailbox"\n'
              '"staff-group","staff@school.example","MailUniversalDistributionGroup"\n'
              '"shared.one","Shared.One@school.example","UserMailbox"\n')
CSV_SAMPLE_2 = ('"Name","PrimarySmtpAddress","RecipientType"\n'
                '"shared.one","shared.one@SCHOOL.example","UserMailbox"\n'
                '"hebakhaled_sec","hebakhaled_sec@school.example","UserMailbox"\n')


def _csv_upload(admin, form):
    return admin.post('/admin/teachers/import-members', data=form, content_type='multipart/form-data',
                      follow_redirects=True).get_data(as_text=True)


def test_members_csv_import_creates_stage_and_links(admin, app):
    from models import TeacherStage
    from services import build_schedule_payload
    # new stage created from the form, from a UTF-8-BOM export
    html = _csv_upload(admin, {'stage_id': 'new', 'new_stage': 'المرحلة الأساسية 1-3',
                               'file': (io.BytesIO(CSV_SAMPLE.encode('utf-8')), 'g.csv')})
    assert 'جديد 4' in html and 'تجاهل 1' in html
    with app.app_context():
        st1 = Stage.query.filter_by(name_ar='المرحلة الأساسية 1-3').one()
        names = sorted(l.teacher.name_ar for l in st1.teacher_links)
        assert names == ["Amira Test", "Esra'a Thyab", 'Rasha AlAyyad', 'Shared One']
        t = Teacher.query.filter_by(email='amira.test@school.example').one()
        assert len(t.code) == 6 and t.stages[0].id == st1.id
        sid1 = st1.id
    # second stage, UTF-16 export, shared teacher matched by email (case-insensitive)
    html = _csv_upload(admin, {'stage_id': 'new', 'new_stage': 'المرحلة الثانوية',
                               'file': (io.BytesIO(CSV_SAMPLE_2.encode('utf-16')), 'g.csv')})
    assert 'جديد 1' in html
    with app.app_context():
        shared = Teacher.query.filter_by(email='shared.one@school.example').one()
        assert sorted(st.name_ar for st in shared.stages) == ['المرحلة الأساسية 1-3', 'المرحلة الثانوية']
        assert Teacher.query.filter_by(email='hebakhaled_sec@school.example').one().name_ar == 'Hebakhaled Sec'
        assert Teacher.query.filter_by(email='rasha.alayyad@school.example').one().name_ar == 'Rasha AlAyyad'
        sid2 = Stage.query.filter_by(name_ar='المرحلة الثانوية').one().id
    # re-import is idempotent
    html = _csv_upload(admin, {'stage_id': sid1, 'file': (io.BytesIO(CSV_SAMPLE.encode()), 'g.csv')})
    assert 'جديد 0' in html
    with app.app_context():
        assert TeacherStage.query.filter_by(stage_id=sid1).count() == 4
    # sync removes members no longer in the group
    html = _csv_upload(admin, {'stage_id': sid1, 'sync': '1',
                               'file': (io.BytesIO(CSV_SAMPLE_2.encode()), 'g.csv')})
    with app.app_context():
        assert sorted(l.teacher.email for l in TeacherStage.query.filter_by(stage_id=sid1)) == \
            ['hebakhaled_sec@school.example', 'shared.one@school.example']
    # empty export (BOM only) and missing stage name are refused
    html = _csv_upload(admin, {'stage_id': sid1, 'file': (io.BytesIO(b'\xef\xbb\xbf'), 'g.csv')})
    assert 'الملف فارغ' in html
    html = _csv_upload(admin, {'stage_id': 'new', 'new_stage': '', 'file': (io.BytesIO(CSV_SAMPLE.encode()), 'g.csv')})
    assert 'اسم المرحلة' in html
    # a new stage with an empty export is still created (import its teachers later)
    html = _csv_upload(admin, {'stage_id': 'new', 'new_stage': 'المرحلة الأساسية 4-6',
                               'file': (io.BytesIO(b'\xef\xbb\xbf'), 'g.csv')})
    assert 'تم إنشاء المرحلة الأساسية 4-6' in html and 'الملف فارغ' in html
    with app.app_context():
        assert Stage.query.filter_by(name_ar='المرحلة الأساسية 4-6').count() == 1

    # schedule payload for a stage-linked teacher: one target per grade, stage name as label
    for kind, dur in [('class', 45), ('class', 45)]:
        admin.post(f'/admin/stages/{sid2}/periods/add', data={'kind': kind, 'duration': dur})
    with app.app_context():
        shared = Teacher.query.filter_by(email='shared.one@school.example').one()
        p = build_schedule_payload('teacher', shared)
        t2 = [t for t in p['targets'] if t['stageId'] == sid2]
        assert len(t2) == 1 and t2[0]['ar'] == 'المرحلة الثانوية'
        assert [x['end'] for x in p['days'][t2[0]['gradeId']]['6']] == ['08:45', '09:30']
    admin.post(f'/admin/stages/{sid2}/grades/add', data={'names': 'العاشر، الحادي عشر'})
    with app.app_context():
        shared = Teacher.query.filter_by(email='shared.one@school.example').one()
        p = build_schedule_payload('teacher', shared)
        assert len([t for t in p['targets'] if t['stageId'] == sid2]) == 2
        tid = shared.id
    # teacher page: stage link can be removed and added
    html = admin.get(f'/admin/teachers/{tid}').get_data(as_text=True)
    assert 'المرحلة الثانوية' in html and 'shared.one@school.example' in html
    with app.app_context():
        link = TeacherStage.query.filter_by(teacher_id=tid, stage_id=sid2).one().id
    admin.post(f'/admin/stage-links/{link}/delete')
    admin.post(f'/admin/teachers/{tid}/stages/add', data={'stage_id': sid2})
    with app.app_context():
        assert TeacherStage.query.filter_by(teacher_id=tid, stage_id=sid2).count() == 1
    # the Excel template of that stage lists the stage-linked teacher with ✓ in "كل المرحلة"
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(admin.get(f'/admin/teachers/template/{sid2}').data))['المعلمون']
    row = next(r for r in ws.iter_rows(min_row=3, values_only=True) if r[2] == 'shared.one@school.example')
    assert row[4] == '✓'
    # renaming to Arabic through the template keeps the same teacher (matched by email)
    data = _fill(admin.get(f'/admin/teachers/template/{sid2}').data, [])
    wb = openpyxl.load_workbook(io.BytesIO(data)); w = wb['المعلمون']
    w['A3'], w['C3'], w['E3'] = 'منار أبو زيد', 'shared.one@school.example', '✓'
    buf = io.BytesIO(); wb.save(buf)
    _upload(admin, sid2, buf.getvalue())
    with app.app_context():
        t = db.session.get(Teacher, tid)
        assert t.name_ar == 'منار أبو زيد' and t.email == 'shared.one@school.example'


def test_schema_migration_adds_email_column(tmp_path):
    import sqlite3
    dbfile = tmp_path / 'old.db'
    con = sqlite3.connect(dbfile)
    con.execute('CREATE TABLE teachers (id INTEGER PRIMARY KEY, name_ar VARCHAR(200) NOT NULL, '
                'name_en VARCHAR(200), active BOOLEAN, code VARCHAR(12) NOT NULL UNIQUE)')
    con.execute("INSERT INTO teachers (name_ar, active, code) VALUES ('قديم', 1, 'ABC234')")
    con.commit(); con.close()
    app = create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{dbfile}'})
    with app.app_context():
        t = Teacher.query.one()
        assert t.name_ar == 'قديم' and t.email is None
        t.email = 'x@y.example'
        db.session.commit()
    create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{dbfile}'})   # idempotent
