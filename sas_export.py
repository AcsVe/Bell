"""«تصدير الحصص لنظام الغياب»: the effective periods of every grade on every school day, as one sheet that
the student-absence system (SAS, 8.iajaward.org) imports in «الهيكل المدرسي ← الحصص ← استيراد».

One row per period: stage, grade, weekday (Python numbering, Monday = 0, and its name), order, kind, the
Arabic and English names, start and end — after every layer (bell program of the day → grade times → stage
day times → stage default), exactly what the bells ring. A stage without grades gives its own day.
Read-only: nothing here changes the bell data.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from models import Stage
from services import WEEK_ORDER, resolve_grade_day, resolve_stage_day, weekday_names

HEAD = ['المرحلة', 'المرحلة بالإنجليزية', 'الصف', 'الصف بالإنجليزية', 'رقم اليوم', 'اليوم', 'الترتيب', 'النوع',
        'الاسم', 'الاسم بالإنجليزية', 'تبدأ', 'تنتهي']


def rows():
    days_ar = weekday_names('ar')
    out = []
    for st in Stage.query.filter(Stage.active.isnot(False)).order_by(Stage.sort_order, Stage.id).all():
        week = [d for d in WEEK_ORDER if d in st.weekday_list]
        owners = [(g, g.name_ar, g.name_en) for g in st.grades] or [(None, '', '')]
        for grade, g_ar, g_en in owners:
            for wd in week:
                day = resolve_grade_day(grade, wd) if grade is not None else resolve_stage_day(st, wd)
                for i, p in enumerate(day, 1):
                    out.append([st.name_ar, st.name_en or '', g_ar, g_en or '', wd, days_ar[wd], i,
                                'استراحة' if p['kind'] == 'break' else 'حصة', p['ar'], p['en'] or '',
                                p['start'], p['end']])
    return out


def export_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = 'الحصص'
    ws.sheet_view.rightToLeft = True
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1B3358')
        c.alignment = Alignment(horizontal='center')
    for r in rows():
        ws.append(r)
    ws.freeze_panes = 'A2'
    for col, w in zip('ABCDEFGHIJKL', (24, 20, 16, 12, 9, 10, 8, 9, 18, 16, 8, 8)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
