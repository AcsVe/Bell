"""Bell programs ⇄ Excel.

Sheet "البرامج" (one row per period, programs listed one after another):
    البرنامج | الترتيب | النوع | التسمية | التسمية بالإنجليزية | من | إلى
Sheet "التوزيع" (one row per grade; each day cell names the program for that day):
    المرحلة | الصف | الصف بالإنجليزية | الأحد | الاثنين | الثلاثاء | الأربعاء | الخميس | الجمعة | السبت
Importing creates missing stages/grades, replaces the listed programs and sets
the grade/day assignments. Times accept 8:00, 08:00, 1:10 (→ 13:10) or Excel times.
"""
import datetime as dt
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from models import (db, BellProgram, BellProgramPeriod, Grade, GradeDayProgram, Stage,
                    StageWeekday)
from services import minutes_between

PROG_SHEET, ASSIGN_SHEET = 'البرامج', 'التوزيع'
PROG_HEAD = ['البرنامج', 'الترتيب', 'النوع', 'التسمية', 'التسمية بالإنجليزية', 'من', 'إلى']
# Python weekday numbers in school-week order (Sunday first)
DAY_COLS = [(6, 'الأحد'), (0, 'الاثنين'), (1, 'الثلاثاء'), (2, 'الأربعاء'), (3, 'الخميس'),
            (4, 'الجمعة'), (5, 'السبت')]
DAY_BY_NAME = {n: wd for wd, n in DAY_COLS}
BREAK_WORDS = ('فرصة', 'الفرصة', 'استراحة', 'الاستراحة', 'break')
# Hours 1–6 written without a.m./p.m. on a school timetable are afternoon.
PM_BEFORE = 7


def _t(v):
    return re.sub(r'\s+', ' ', str(v if v is not None else '')).strip()


def parse_time(v):
    """→ 'HH:MM' or None."""
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        h, m = v.hour, v.minute
    elif isinstance(v, (int, float)) and 0 <= v < 1:          # Excel fraction of a day
        total = round(v * 24 * 60)
        h, m = divmod(total, 60)
    else:
        mt = re.fullmatch(r'\s*(\d{1,2})\s*[:.٫]\s*(\d{2})\s*', _t(v))
        if not mt:
            return None
        h, m = int(mt.group(1)), int(mt.group(2))
    if 1 <= h < PM_BEFORE:
        h += 12
    if not (0 <= h < 24 and 0 <= m < 60):
        return None
    return f'{h:02d}:{m:02d}'


def _kind(kind_cell, label):
    k = _t(kind_cell).lower()
    if k in ('استراحة', 'فرصة', 'break', 'الفرصة', 'الاستراحة'):
        return 'break'
    if k in ('حصة', 'class', 'period'):
        return 'class'
    return 'break' if any(w in _t(label).lower() for w in BREAK_WORDS) else 'class'


def _style_head(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='12305E')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = 'A2'


def export_programs():
    wb = Workbook()
    ws = wb.active
    ws.title = PROG_SHEET
    ws.append(PROG_HEAD)
    for prog in BellProgram.query.order_by(BellProgram.sort_order, BellProgram.id):
        for p in prog.periods:
            ws.append([prog.name, p.number, 'استراحة' if p.kind == 'break' else 'حصة',
                       p.label_ar, p.label_en or '', p.start_time, p.end_time])
    for col, w in zip('ABCDEFG', (30, 9, 10, 20, 16, 9, 9)):
        ws.column_dimensions[col].width = w
    _style_head(ws)

    wa = wb.create_sheet(ASSIGN_SHEET)
    wa.append(['المرحلة', 'الصف', 'الصف بالإنجليزية'] + [n for _, n in DAY_COLS])
    for st in Stage.query.order_by(Stage.sort_order, Stage.id):
        for g in st.grades:
            progs = {a.weekday: a.program.name for a in g.day_programs}
            wa.append([st.name_ar, g.name_ar, g.name_en or ''] + [progs.get(wd, '') for wd, _ in DAY_COLS])
    for col, w in zip('ABCDEFGHIJ', (26, 18, 14, 22, 22, 22, 22, 22, 12, 12)):
        wa.column_dimensions[col].width = w
    _style_head(wa)

    info = wb.create_sheet('تعليمات')
    info.sheet_view.rightToLeft = True
    info.column_dimensions['A'].width = 110
    for line in [
        'برامج الجرس',
        '',
        '• ورقة «البرامج»: كل سطر حصة أو فرصة. اكتب اسم البرنامج في كل سطر من سطوره، والتوقيت بصيغة 8:00 أو 1:10 (الساعات 1–6 تُعتبر بعد الظهر).',
        '• النوع: حصة أو استراحة (إن تُرك فارغاً يُستنتج من التسمية: «الفرصة/الاستراحة» = استراحة).',
        '• التسمية هي نص التنبيه: «انتهت الحصة الأولى - بداية الحصة الثانية».',
        '• ورقة «التوزيع»: كل سطر صف. اكتب في خانة كل يوم اسم البرنامج الذي يتبعه الصف ذلك اليوم، واتركها فارغة إن لم يكن دوام.',
        '• المرحلة والصف يُنشآن تلقائياً إن لم يكونا موجودين. لنقل صف إلى مرحلة أخرى غيّر اسم المرحلة في سطره.',
        '• رفع الملف يستبدل البرامج المذكورة فيه ويحدّث توزيع الصفوف المذكورة فقط.',
    ]:
        info.append([line])
    info['A1'].font = Font(bold=True, size=14)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def import_programs(fileobj):
    """Returns a summary; raises ValueError (nothing saved) on any error."""
    try:
        wb = load_workbook(fileobj, data_only=True)
    except Exception:
        raise ValueError('الملف ليس ملف Excel صالحاً (xlsx).')
    if PROG_SHEET not in wb.sheetnames:
        raise ValueError(f'لم أجد ورقة «{PROG_SHEET}» في الملف.')

    errors, warnings = [], []
    programs = {}                       # name → list of period dicts (in file order)
    rows = list(wb[PROG_SHEET].iter_rows(values_only=True))
    for i, r in enumerate(rows[1:], start=2):
        r = list(r) + [None] * 7
        name, label = _t(r[0]), _t(r[3])
        if not name and not label:
            continue
        if not name or not label:
            errors.append(f'البرامج سطر {i}: اسم البرنامج والتسمية مطلوبان')
            continue
        start, end = parse_time(r[5]), parse_time(r[6])
        if not start or not end or start >= end:
            errors.append(f'البرامج سطر {i} ({name} / {label}): توقيت غير صالح «{_t(r[5])} – {_t(r[6])}»')
            continue
        programs.setdefault(name, []).append({
            'order': r[1] if isinstance(r[1], (int, float)) else len(programs.get(name, [])) + 1,
            'kind': _kind(r[2], label), 'label_ar': label, 'label_en': _t(r[4]) or None,
            'start': start, 'end': end})

    for name, ps in programs.items():
        ps.sort(key=lambda p: (p['order'], p['start']))
        for a, b in zip(ps, ps[1:]):
            if b['start'] < a['end']:
                errors.append(f'{name}: «{a["label_ar"]}» و«{b["label_ar"]}» متداخلتان')
            elif b['start'] > a['end']:
                warnings.append(f'{name}: فجوة {minutes_between(a["end"], b["start"])} د بين «{a["label_ar"]}» و«{b["label_ar"]}»')

    assignments = []
    if ASSIGN_SHEET in wb.sheetnames:
        arows = list(wb[ASSIGN_SHEET].iter_rows(values_only=True))
        head = [_t(h) for h in (arows[0] if arows else [])]
        day_idx = {DAY_BY_NAME[h]: c for c, h in enumerate(head) if h in DAY_BY_NAME}
        existing_names = {p.name for p in BellProgram.query.all()}
        for i, r in enumerate(arows[1:], start=2):
            r = list(r) + [None] * len(head)
            stage_name, grade_name = _t(r[0]), _t(r[1])
            if not stage_name and not grade_name:
                continue
            if not stage_name or not grade_name:
                errors.append(f'التوزيع سطر {i}: المرحلة والصف مطلوبان')
                continue
            days = {}
            for wd, c in day_idx.items():
                pn = _t(r[c])
                if pn and pn not in programs and pn not in existing_names:
                    errors.append(f'التوزيع سطر {i} ({grade_name}): البرنامج «{pn}» غير موجود')
                days[wd] = pn or None
            assignments.append({'stage': stage_name, 'grade': grade_name,
                                'grade_en': _t(r[2]) or None, 'days': days})

    if errors:
        more = f' … و{len(errors) - 8} أخطاء أخرى' if len(errors) > 8 else ''
        raise ValueError('لم يُحفظ شيء: ' + ' | '.join(errors[:8]) + more)

    # ── save (all or nothing) ──
    summary = {'programs': 0, 'stages': 0, 'grades': 0, 'assigned': 0, 'warnings': warnings}
    by_name = {p.name: p for p in BellProgram.query.all()}
    for order, (name, ps) in enumerate(programs.items()):
        prog = by_name.get(name)
        if prog is None:
            prog = BellProgram(name=name, sort_order=BellProgram.query.count() + order)
            db.session.add(prog)
            by_name[name] = prog
        else:
            for old in list(prog.periods):
                db.session.delete(old)
            db.session.flush()
        for n, p in enumerate(ps, start=1):
            prog.periods.append(BellProgramPeriod(number=n, kind=p['kind'], label_ar=p['label_ar'],
                                                  label_en=p['label_en'], start_time=p['start'],
                                                  end_time=p['end']))
        summary['programs'] += 1
    db.session.flush()

    stages = {s.name_ar: s for s in Stage.query.all()}
    for a in assignments:
        st = stages.get(a['stage'])
        if st is None:
            st = Stage(name_ar=a['stage'], sort_order=len(stages))
            db.session.add(st)
            db.session.flush()
            stages[a['stage']] = st
            summary['stages'] += 1
        g = next((g for g in st.grades if g.name_ar == a['grade']), None)
        if g is None:
            # A grade moved here from another stage keeps its sections and codes.
            g = Grade.query.filter_by(name_ar=a['grade']).first()
            if g is not None and g.stage_id != st.id:
                g.stage_id = st.id
            elif g is None:
                g = Grade(stage_id=st.id, name_ar=a['grade'], sort_order=len(st.grades))
                db.session.add(g)
                summary['grades'] += 1
            db.session.flush()
            db.session.expire(st, ['grades'])
        if a['grade_en']:
            g.name_en = a['grade_en']
        current = {x.weekday: x for x in g.day_programs}
        for wd, pn in a['days'].items():
            if pn:
                prog = by_name[pn]
                if wd in current:
                    current[wd].program_id = prog.id
                else:
                    db.session.add(GradeDayProgram(grade_id=g.id, weekday=wd, program_id=prog.id))
                summary['assigned'] += 1
                if wd not in st.weekday_list:
                    db.session.add(StageWeekday(stage_id=st.id, weekday=wd))
                    db.session.flush()
                    db.session.expire(st, ['weekdays'])
            elif wd in current:
                db.session.delete(current[wd])
    db.session.commit()
    return summary
